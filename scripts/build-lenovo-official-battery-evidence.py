#!/usr/bin/env python3
"""Match RB Lenovo title candidates to Lenovo's official battery registry.

The command produces evidence only.  It never mutates products, publication,
SEO, media, prices, compatibility, or canonical identity.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl


SOURCE_URL = (
    "https://www.lenovo.com/content/dam/lenovo/site-design/esg-document-library/"
    "global/battery-msds/BatteryInformationFinder.xlsx"
)
SHEET_NAME = "BatteryData"
REQUIRED_CANDIDATE_COLUMNS = {
    "product_external_id",
    "name",
    "brand_candidate",
    "part_number_candidates",
    "primary_part_number_key",
    "voltage_v_candidates",
    "capacity_mah_candidates",
    "energy_wh_candidates",
    "identity_risk",
    "safe_to_apply",
}
REQUIRED_OFFICIAL_HEADERS = {
    "ASSM PN",
    "FRU PN",
    "Battery Model",
    "WH Rating",
    "Cell Voltage (VDC)",
    "Battery Voltage (VDC)",
    "BatterySupplier",
    "SDS Link",
    "UN38.3 Link",
}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_part(value: Any) -> str:
    return re.sub(r"[^A-Z0-9+]", "", str(value or "").upper())


def official_part_keys(value: Any) -> list[str]:
    raw = str(value or "").strip()
    if not raw or raw.upper() in {"N/A", "NA", "NONE", "0"}:
        return []
    result: list[str] = []
    for token in re.findall(r"[A-Z0-9][A-Z0-9._+/-]{2,30}", raw.upper()):
        key = normalize_part(token)
        if len(key) >= 3 and any(character.isdigit() for character in key):
            result.append(key)
    return list(dict.fromkeys(result))


def number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        parsed = float(str(value).replace(",", "."))
    except ValueError:
        return None
    return parsed if math.isfinite(parsed) else None


def display_number(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:.4f}".rstrip("0").rstrip(".")


def first_candidate(value: str) -> float | None:
    first = (value or "").split("|", 1)[0].strip()
    return number(first)


def comparison(candidate: float | None, official: float | None, tolerance: float) -> str:
    if official is None:
        return "official_value_missing"
    if candidate is None:
        return "title_value_missing"
    return "match" if abs(candidate - official) <= tolerance else "conflict"


def load_candidates(path: Path, expected_records: int) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        missing = REQUIRED_CANDIDATE_COLUMNS - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"missing candidate columns: {', '.join(sorted(missing))}")
        rows = list(reader)
    if len(rows) != expected_records:
        raise ValueError(f"candidate count mismatch: expected {expected_records}, got {len(rows)}")
    ids = [row["product_external_id"] for row in rows]
    if len(ids) != len(set(ids)):
        raise ValueError("duplicate product_external_id in candidate registry")
    return rows


def load_official_rows(path: Path, expected_rows: int) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"official workbook is missing {SHEET_NAME}")
    sheet = workbook[SHEET_NAME]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(iterator)]
    missing = REQUIRED_OFFICIAL_HEADERS - set(headers)
    if missing:
        raise ValueError(f"missing official columns: {', '.join(sorted(missing))}")
    rows = [dict(zip(headers, values)) for values in iterator]
    workbook.close()
    if len(rows) != expected_rows:
        raise ValueError(f"official row count mismatch: expected {expected_rows}, got {len(rows)}")
    return rows


def official_signature(row: dict[str, Any]) -> tuple[str, ...]:
    return (
        str(row.get("ASSM PN") or ""),
        str(row.get("FRU PN") or ""),
        str(row.get("Battery Model") or ""),
        display_number(number(row.get("WH Rating"))),
        display_number(number(row.get("Cell Voltage (VDC)"))),
        display_number(number(row.get("Battery Voltage (VDC)"))),
        str(row.get("BatterySupplier") or ""),
        str(row.get("SDS Link") or ""),
        str(row.get("UN38.3 Link") or ""),
    )


def build(
    candidates_path: Path,
    official_path: Path,
    output_path: Path,
    summary_path: Path,
    expected_candidate_records: int,
    expected_selected_records: int,
    expected_official_rows: int,
    expected_official_sha256: str,
) -> dict[str, object]:
    actual_official_sha256 = file_sha256(official_path)
    if actual_official_sha256.lower() != expected_official_sha256.lower():
        raise ValueError(
            "official workbook SHA-256 mismatch: "
            f"expected {expected_official_sha256}, got {actual_official_sha256}"
        )
    candidates = load_candidates(candidates_path, expected_candidate_records)
    selected = [
        row for row in candidates
        if row["brand_candidate"] == "Lenovo"
        and row["identity_risk"] == "single_candidate_needs_primary_source"
        and row["safe_to_apply"].casefold() == "false"
    ]
    if len(selected) != expected_selected_records:
        raise ValueError(
            f"selected Lenovo count mismatch: expected {expected_selected_records}, got {len(selected)}"
        )

    official_rows = load_official_rows(official_path, expected_official_rows)
    index: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in official_rows:
        for key in set(official_part_keys(row.get("ASSM PN")) + official_part_keys(row.get("FRU PN"))):
            index[key].append(row)

    evidence: list[dict[str, str]] = []
    for candidate in selected:
        key = candidate["primary_part_number_key"]
        matches_by_signature = {
            official_signature(row): row for row in index.get(key, [])
        }
        matches = list(matches_by_signature.values())
        official = matches[0] if len(matches) == 1 else None
        official_voltage = number(official.get("Battery Voltage (VDC)")) if official else None
        official_wh = number(official.get("WH Rating")) if official else None
        derived_mah = (
            official_wh * 1000 / official_voltage
            if official_wh is not None and official_voltage not in (None, 0)
            else None
        )
        if not matches:
            status = "not_found_official_registry"
        elif len(matches) > 1:
            status = "official_exact_conflict_hold"
        else:
            status = "official_exact_unique"

        evidence.append({
            "product_external_id": candidate["product_external_id"],
            "name": candidate["name"],
            "candidate_part_number": candidate["part_number_candidates"],
            "candidate_part_number_key": key,
            "official_match_status": status,
            "official_match_count": str(len(matches)),
            "official_asm_pn": str(official.get("ASSM PN") or "") if official else "",
            "official_fru_pn": str(official.get("FRU PN") or "") if official else "",
            "official_battery_model": str(official.get("Battery Model") or "") if official else "",
            "official_supplier": str(official.get("BatterySupplier") or "") if official else "",
            "official_cell_voltage_v": display_number(number(official.get("Cell Voltage (VDC)"))) if official else "",
            "official_battery_voltage_v": display_number(official_voltage),
            "official_energy_wh": display_number(official_wh),
            "derived_capacity_mah_audit_only": display_number(derived_mah),
            "title_voltage_v": candidate["voltage_v_candidates"],
            "title_capacity_mah": candidate["capacity_mah_candidates"],
            "title_energy_wh": candidate["energy_wh_candidates"],
            "voltage_comparison": comparison(first_candidate(candidate["voltage_v_candidates"]), official_voltage, 0.05) if official else "not_compared",
            "energy_comparison": comparison(first_candidate(candidate["energy_wh_candidates"]), official_wh, 0.1) if official else "not_compared",
            "official_sds_url": str(official.get("SDS Link") or "") if official else "",
            "official_un38_3_url": str(official.get("UN38.3 Link") or "") if official else "",
            "official_registry_url": SOURCE_URL,
            "evidence_scope": "exact_oem_part_regulatory_data",
            "compatibility_verified": "false",
            "image_verified": "false",
            "capacity_publication_allowed": "false",
            "safe_to_apply": "false",
            "review_status": "official_facts_staged" if status == "official_exact_unique" else "hold",
        })

    evidence.sort(key=lambda row: row["candidate_part_number_key"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(evidence[0]))
        writer.writeheader()
        writer.writerows(evidence)

    status_counts = Counter(row["official_match_status"] for row in evidence)
    voltage_counts = Counter(row["voltage_comparison"] for row in evidence)
    energy_counts = Counter(row["energy_comparison"] for row in evidence)
    summary: dict[str, object] = {
        "candidate_source_path": str(candidates_path),
        "candidate_source_sha256": file_sha256(candidates_path),
        "official_source_path": str(official_path),
        "official_source_url": SOURCE_URL,
        "official_source_sha256": actual_official_sha256,
        "official_battery_rows": len(official_rows),
        "selected_lenovo_records": len(evidence),
        "official_match_status_counts": dict(sorted(status_counts.items())),
        "voltage_comparison_counts": dict(sorted(voltage_counts.items())),
        "energy_comparison_counts": dict(sorted(energy_counts.items())),
        "official_unique_with_voltage": sum(
            row["official_match_status"] == "official_exact_unique"
            and bool(row["official_battery_voltage_v"])
            for row in evidence
        ),
        "official_unique_with_energy": sum(
            row["official_match_status"] == "official_exact_unique"
            and bool(row["official_energy_wh"])
            for row in evidence
        ),
        "automatic_database_mutations": 0,
        "canonical_identity_changes": 0,
        "compatibility_changes": 0,
        "image_changes": 0,
        "published_fact_changes": 0,
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidates", type=Path, required=True)
    parser.add_argument("--official-workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    parser.add_argument("--expected-candidate-records", type=int, required=True)
    parser.add_argument("--expected-selected-records", type=int, required=True)
    parser.add_argument("--expected-official-rows", type=int, required=True)
    parser.add_argument("--expected-official-sha256", required=True)
    args = parser.parse_args()
    summary = build(
        args.candidates,
        args.official_workbook,
        args.output,
        args.summary,
        args.expected_candidate_records,
        args.expected_selected_records,
        args.expected_official_rows,
        args.expected_official_sha256,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
