#!/usr/bin/env python3
"""Build the full site-catalog mapping file for the owner to align with 1C.

Read-only join of already-generated, gitignored extractor artifacts:
  docs/audits/generated/bitrix-b2b-catalog-sections.csv   (219 groups)
  docs/audits/generated/bitrix-b2b-catalog-products.csv   (2856 products)
  docs/audits/generated/rb-import-manifest-draft.csv       (price/brand/mpn candidates)

Emits an owner-facing workbook + CSVs under docs/imports/ with EVERY group and
EVERY product plus empty 1C columns to fill. Nothing is invented: price/stock are
carried only where the backup had them and are explicitly dated + flagged; brand
and MPN are carried only as *candidates*. This produces no SKU/1C identity — that
is exactly what the owner supplies from 1C.

Usage (from repo root):
  python scripts/build-site-catalog-for-1c-mapping.py
"""
from __future__ import annotations

import csv
import os
import sys

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GEN = os.path.join(REPO, "docs", "audits", "generated")
OUT = os.path.join(REPO, "docs", "imports")
SITE = "https://microchips.by"
BACKUP_DATE = "2026-06-23"


def read_csv(path: str) -> list[dict]:
    if not os.path.exists(path):
        sys.exit(f"Required source not found (run the extractor first): {path}")
    # utf-8-sig strips the BOM the PowerShell extractor writes.
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def product_url(row: dict) -> str:
    path = (row.get("legacy_url_candidate") or "").strip()
    return f"{SITE}{path}" if path.startswith("/") else ""


def section_url(path: str) -> str:
    path = (path or "").strip().strip("/")
    return f"{SITE}/catalog/{path}/" if path else ""


def build_rows() -> tuple[list[list], list[list]]:
    sections = read_csv(os.path.join(GEN, "bitrix-b2b-catalog-sections.csv"))
    products = read_csv(os.path.join(GEN, "bitrix-b2b-catalog-products.csv"))
    manifest = {r["legacy_element_id"]: r for r in read_csv(os.path.join(GEN, "rb-import-manifest-draft.csv"))}

    cat_rows = [[
        "Bitrix ID раздела", "Родитель ID", "Название", "Путь на сайте",
        "Ссылка на сайт", "🟥 1С-группа (заполнить)",
    ]]
    for s in sections:
        path = s.get("source_section_path", "")
        cat_rows.append([
            s.get("legacy_section_id", ""), s.get("parent_section_id", ""),
            s.get("name", ""), path, section_url(path), "",
        ])

    prod_rows = [[
        "Bitrix ID", "Название", "Раздел (путь)", "Ссылка на сайт",
        "Ёмкость (из названия)", "Напряжение (из названия)", "Технология (из названия)",
        "Бренд? (кандидат)", "MPN? (кандидат)",
        f"Цена BYN (бэкап {BACKUP_DATE})", "Статус цены", "В наличии (бэкап)",
        "🟥 1С-код (заполнить)", "🟥 SKU (заполнить)",
    ]]
    for p in products:
        pid = p.get("legacy_element_id", "")
        m = manifest.get(pid, {})
        prod_rows.append([
            pid, p.get("name", ""), p.get("primary_section_path", ""), product_url(p),
            p.get("capacity_from_name", ""), p.get("voltage_from_name", ""), p.get("technology_from_name", ""),
            m.get("brand_candidate", ""), m.get("mpn_candidate_from_name", ""),
            m.get("price_byn", ""), m.get("price_status", "нет данных в бэкапе"), m.get("in_stock", ""),
            "", "",
        ])
    return cat_rows, prod_rows


def write_csv(path: str, rows: list[list]) -> None:
    # UTF-8 with BOM so Excel/Google Sheets read Cyrillic correctly on open.
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh).writerows(rows)


def write_xlsx(path: str, cat_rows: list[list], prod_rows: list[list]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return False

    wb = Workbook()
    warn = (f"Данные сайта из бэкапа Bitrix ({BACKUP_DATE}) — НЕ живые. Цена/наличие только где были в бэкапе; "
            "бренд/MPN — кандидаты из названия, не подтверждены. SKU/1С-код заполняет владелец из 1С.")
    red = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
    head = Font(bold=True)

    def fill(ws, rows):
        ws.append([warn] + [""] * (len(rows[0]) - 1))
        ws.append(rows[0])
        for r in rows[1:]:
            ws.append(r)
        for cell in ws[2]:
            cell.font = head
            if isinstance(cell.value, str) and cell.value.startswith("🟥"):
                cell.fill = red
        ws.freeze_panes = "A3"

    ws1 = wb.active
    ws1.title = "Товары"
    fill(ws1, prod_rows)
    ws2 = wb.create_sheet("Категории")
    fill(ws2, cat_rows)
    wb.save(path)
    return True


def main() -> None:
    os.makedirs(OUT, exist_ok=True)
    cat_rows, prod_rows = build_rows()

    write_csv(os.path.join(OUT, "site-catalog-full-categories.csv"), cat_rows)
    write_csv(os.path.join(OUT, "site-catalog-full-products.csv"), prod_rows)
    xlsx_path = os.path.join(OUT, "site-catalog-full-for-1c-mapping.xlsx")
    has_xlsx = write_xlsx(xlsx_path, cat_rows, prod_rows)

    print(f"Категорий (групп): {len(cat_rows) - 1}")
    print(f"Товаров: {len(prod_rows) - 1}")
    priced = sum(1 for r in prod_rows[1:] if r[9])
    print(f"  из них с ценой из бэкапа: {priced}")
    print(f"CSV: {os.path.join(OUT, 'site-catalog-full-products.csv')}")
    print(f"CSV: {os.path.join(OUT, 'site-catalog-full-categories.csv')}")
    print(f"XLSX: {xlsx_path}" if has_xlsx else "XLSX: пропущен (openpyxl не установлен) — используйте CSV")


if __name__ == "__main__":
    main()
