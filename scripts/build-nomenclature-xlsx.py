# -*- coding: utf-8 -*-
"""Build a Google-Sheets-ready .xlsx from the RB nomenclature findings.

Two sheets:
  «К исправлению» — one row per finding with reviewer columns (status, correct
                    value, who/when), clickable legacy-card links and severity
                    colour coding.
  «Сводка»        — counts per error class plus a short how-to note.

Used two ways:
  * imported by ``audit-rb-nomenclature.py`` (``build_xlsx(findings, out)``) so
    the workbook is refreshed on every audit run;
  * run standalone to rebuild the workbook from an existing findings CSV
    without re-running the whole audit:  ``python scripts/build-nomenclature-xlsx.py``.

Depends on ``openpyxl``.
"""
import csv
import os
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = 'https://microchips.by'
FONT = 'Arial'

# Compact, xlsx-friendly titles (mirror the audit's CLASS_INFO, shortened).
CLASS_TITLE = {
    'E1_voltage_mismatch': 'Напряжение: название ≠ база',
    'E2_capacity_mismatch': 'Ёмкость: название ≠ база',
    'E3_technology_family_mismatch': 'Химия: грубое расхождение',
    'E4_technology_subtype_mismatch': 'Технология: AGM ≠ GEL',
    'E5_identity_collision_contradictory': 'Один артикул — разные товары',
    'E6_identity_collision_duplicate': 'Один артикул — вероятный дубль',
    'E7_duplicate_name': 'Одинаковые названия',
    'E8_mixed_script_token': 'Латиница + кириллица в коде',
    'E9_missing_all_specs': 'Нет ключевых характеристик',
    'E10_implausible_voltage': 'Подозрительное напряжение',
}
SEV_RU = {'high': 'Высокая', 'medium': 'Средняя', 'low': 'Низкая'}
SEV_ORDER = {'high': 0, 'medium': 1, 'low': 2}
COLLISION_PREFIXES = ('E5', 'E6', 'E7')

# Column layout for the «К исправлению» sheet.
HEADERS = ['Важность', 'Класс проблемы', 'ID', 'Товар', 'Ссылка на карточку',
           'Поле', 'В названии', 'В базе', 'Что не так', 'Рекомендация',
           'Статус исправления', 'Правильное значение', 'Кто / когда исправил']
WIDTHS = [11, 26, 8, 42, 34, 14, 14, 30, 34, 40, 18, 20, 20]
WRAP_COLS = {4, 8, 9, 10}          # 1-based columns that wrap text
REVIEWER_COLS = (11, 12, 13)       # editable columns highlighted for the owner


def _sev_fill(sev):
    return PatternFill('solid', fgColor={
        'high': 'F4CCCC', 'medium': 'FCE5CD', 'low': 'FFF2CC',
    }.get(sev, 'FFFFFF'))


def build_xlsx(rows, out_path, focus_count=None):
    """Write the two-sheet workbook to ``out_path``.

    ``rows`` is a list of finding dicts (same shape as the findings CSV /
    the audit's in-memory findings). Returns ``(out_path, row_count)``.
    """
    rows = sorted(rows, key=lambda r: (SEV_ORDER.get(r['severity'], 9),
                                       r['error_class'], r['name']))

    wb = Workbook()
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill('solid', fgColor='38761D')
    hint_fill = PatternFill('solid', fgColor='FFFDE7')

    # ---------------- Sheet 1: К исправлению ----------------
    ws = wb.active
    ws.title = 'К исправлению'
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF', size=11)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    for r in rows:
        url = (r.get('legacy_url_candidate') or '').strip()
        full = (BASE_URL + url) if url.startswith('/') else ''
        is_collision = r['error_class'].startswith(COLLISION_PREFIXES)
        in_base = (r['value_in_property'].replace('||', ' ↔ ')
                   if is_collision else r['value_in_property'])
        ws.append([
            SEV_RU.get(r['severity'], r['severity']),
            CLASS_TITLE.get(r['error_class'], r['error_class']),
            r['legacy_element_id'],
            r['name'],
            full,
            r['field'],
            r['value_in_name'],
            in_base,
            r['note'],
            r['suggested_action'],
            '', '', '',
        ])
        rr = ws.max_row
        ws.cell(row=rr, column=1).fill = _sev_fill(r['severity'])
        if full:
            link = ws.cell(row=rr, column=5)
            link.hyperlink = full
            link.value = url
            link.font = Font(name=FONT, color='1155CC', underline='single', size=10)
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=rr, column=c)
            if c != 5:
                cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in WRAP_COLS))
            cell.border = border
        for c in REVIEWER_COLS:
            ws.cell(row=rr, column=c).fill = hint_fill

    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}{ws.max_row}'

    # ---------------- Sheet 2: Сводка ----------------
    ws2 = wb.create_sheet('Сводка')
    ws2.append(['Аудит номенклатуры microchips.by — нестыковки для исправления'])
    ws2['A1'].font = Font(name=FONT, bold=True, size=14)
    affected = len({r['legacy_element_id'] for r in rows})
    if focus_count is not None:
        ws2.append([f'Проверено позиций фокуса: {focus_count}'])
    ws2.append([f'Товаров с проблемами: {affected}'])
    ws2.append([f'Всего находок: {len(rows)}'])
    ws2.append([])

    head_row = ws2.max_row + 1
    ws2.append(['Класс проблемы', 'Важность', 'Находок'])
    for c in range(1, 4):
        cell = ws2.cell(row=head_row, column=c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = border

    cnt = Counter(r['error_class'] for r in rows)
    sev_of = {r['error_class']: r['severity'] for r in rows}
    for cls, n in sorted(cnt.items(), key=lambda kv: (SEV_ORDER[sev_of[kv[0]]], kv[0])):
        ws2.append([CLASS_TITLE.get(cls, cls), SEV_RU.get(sev_of[cls], sev_of[cls]), n])
        rr = ws2.max_row
        ws2.cell(row=rr, column=2).fill = _sev_fill(sev_of[cls])
        for c in range(1, 4):
            ws2.cell(row=rr, column=c).font = Font(name=FONT, size=10)
            ws2.cell(row=rr, column=c).border = border

    ws2.append([])
    note = ('Источник не изменялся. После правок в 1С и на старом сайте — снять свежую '
            'выгрузку и перезапустить scripts/audit-rb-nomenclature.py: отчёт покажет, '
            'что закрыто. «Один артикул — разные товары»: дать разные артикулы или '
            'варианты (например, с/без электролита, AGM/GEL).')
    ws2.append([note])
    note_row = ws2.max_row
    ws2.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    ws2.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=3)
    ws2.row_dimensions[note_row].height = 80
    ws2.column_dimensions['A'].width = 40
    ws2.column_dimensions['B'].width = 14
    ws2.column_dimensions['C'].width = 12

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path, len(rows)


def _repo_paths():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    csv_path = os.path.join(root, 'docs', 'audits', 'generated', 'rb-nomenclature-findings.csv')
    out_path = os.path.join(root, 'docs', 'imports', 'microchips-nomenclature-fixes.xlsx')
    return csv_path, out_path


def main():
    csv_path, out_path = _repo_paths()
    if not os.path.exists(csv_path):
        sys.exit(f'Missing findings CSV: {csv_path}. '
                 'Run scripts/audit-rb-nomenclature.py first.')
    with open(csv_path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    path, n = build_xlsx(rows, out_path)
    print('saved', path, f'({n} rows)')


if __name__ == '__main__':
    main()
